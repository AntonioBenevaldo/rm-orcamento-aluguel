from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt


def require_text(paragraph, expected: str, number: int) -> None:
    actual = paragraph.text
    if actual != expected:
        raise RuntimeError(
            f"Paragraph {number} changed unexpectedly. Expected {expected!r}, got {actual!r}."
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()

    doc = Document(args.input)

    # Keep the cover/summary page self-contained; the introduction was orphaned
    # as the final line of the first page.
    doc.paragraphs[35].paragraph_format.page_break_before = True

    # Remove indentation made with literal spaces.
    p178 = doc.paragraphs[177]
    require_text(
        p178,
        "       Sem crianças: aplicar 5% sobre o subtotal do aluguel.",
        178,
    )
    p178.runs[0].text = ""
    old_ppr = p178._p.pPr
    if old_ppr is not None:
        p178._p.remove(old_ppr)
    p178._p.insert(0, deepcopy(doc.paragraphs[176]._p.pPr))
    for paragraph in (doc.paragraphs[175], doc.paragraphs[176], p178):
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(2)

    # Restore the missing space after the colon.
    p184 = doc.paragraphs[183]
    require_text(
        p184,
        "Validação importante:A interface deve aceitar zero vaga ou pelo menos duas vagas",
        184,
    )
    p184.runs[2].text = "importante: A"

    # Remove a trailing space at the manual paragraph break.
    p207 = doc.paragraphs[206]
    require_text(
        p207,
        "Importante: Para comprovar TDD no repositório, é recomendável preservar commits ",
        207,
    )
    p207.runs[-1].text = ""

    # Correct two words that were truncated/altered in the continuation line.
    p208 = doc.paragraphs[207]
    require_text(
        p208,
        "ue mostrem o teste falhando antes da implementação e, depois, o mesmo teste provado.",
        208,
    )
    p208.runs[0].text = "que"
    p208.runs[-1].text = "aprovado."

    # Normalize spacing and use the actual ignored folder/file names.
    p294 = doc.paragraphs[293]
    require_text(
        p294,
        "gitignore para excluir .venv,   pycache  , bancos temporários e arquivos locais desnecessários.",
        294,
    )
    p294.runs[0].text = ".gitignore"
    p294.runs[8].text = ""
    p294.runs[9].text = "__pycache__"
    p294.runs[10].text = ""

    # Remove literal spaces placed before the GitHub hyperlink.
    p299 = doc.paragraphs[298]
    require_text(
        p299,
        "   https://github.com/AntonioBenevaldo/rm-orcamento-aluguel",
        299,
    )
    p299.runs[0].text = ""

    # Replace tab characters inside the bibliographic reference with normal spaces.
    p314 = doc.paragraphs[313]
    require_text(
        p314,
        "THE\tPANDAS\tDEVELOPMENT\tTEAM.\tpandas\tDocumentation.\tDisponível em: https://pandas.pydata.org/docs/. Acesso em: 19 jul. 2026.",
        314,
    )
    for run in p314.runs:
        if run.text == "\t":
            run.text = " "

    # A handful of imported paragraphs had nearly double line spacing, which
    # produced large gaps inside short instructions and the video script.
    for number in (90, 91, 92, 132, 133, 144, 147, 211, 212, 217, 219, 243, 304, 305):
        doc.paragraphs[number - 1].paragraph_format.line_spacing = 1.3

    # The apartment formula already appears in the bordered formula block. The
    # immediately repeated plain-text copy was accidental.
    p175 = doc.paragraphs[174]
    require_text(
        p175,
        "aluguel = (700 + adicional_quarto + adicional_garagem) - desconto",
        175,
    )
    p175._element.getparent().remove(p175._element)

    # Replace the manually typed dot leaders with real right-aligned tab leaders
    # and synchronize the section names/page numbers with the final pagination.
    toc = (
        ("1. Introdução e objetivos", 2),
        ("2. Modelagem do problema e requisitos", 2),
        ("3. Requisitos funcionais e não funcionais", 3),
        ("4. Modelagem estática e modelagem dinâmica", 5),
        ("5. Modelagem de dados", 6),
        ("6. Arquitetura", 7),
        ("7. Pensamento algorítmico", 8),
        ("8. Fluxograma da aplicação", 9),
        ("9. Regras de negócio e cálculos", 10),
        ("10. Orientação a objetos", 11),
        ("11. TDD - Test-Driven Development", 12),
        ("12. Implementação tecnológica", 14),
        ("13. Integridade SQLite", 14),
        ("14. Geração do CSV com 12 meses", 16),
        ("15. Testes e validação", 16),
        ("16. Execução do projeto", 17),
        ("17. Entregáveis", 18),
        ("18. Conclusão", 18),
        ("19. Referências", 19),
    )
    for paragraph, (label, page) in zip(doc.paragraphs[15:34], toc):
        paragraph.text = f"{label}\t{page}"
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            Inches(4.75), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output.resolve())


if __name__ == "__main__":
    main()
